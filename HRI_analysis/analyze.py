"""
HRI robot trust study: automatic analysis (trust is the main outcome).

Usage:
    python analyze.py HRI_data_entry.xlsx
    python analyze.py HRI_data_entry.xlsx --out results

Reads three sheets typed by hand:
    Data   participant, level, question, source, advice, correct, chosen, time_s   (one row per question)
    Trust  participant, level, trust_robot, trust_screen, rated_first             (one row per participant and level)
    Final  participant, more_trusted, reason                                       (one row per participant)

It checks the typing (typos stop the program and name the sheet and Excel row), computes everything
else, and writes into the output folder:
    summary.txt        plain-language results, trust first
    HRI_results.xlsx   all tables, one sheet each
    *.png              charts

Needs: pip install pandas openpyxl scipy matplotlib   (no statsmodels needed)
"""
import argparse
import os
import sys
import warnings

import numpy as np
import pandas as pd
from scipy import stats

warnings.filterwarnings("ignore")

# ==========================================================================
# CONFIG: change these if the design changes. Set a DESIGN value to None to switch that check off.
# ==========================================================================
LEVELS = ["easy", "medium", "hard"]
SOURCES = ["robot", "screen"]
BOWLS = ["red", "blue"]
FINAL_CHOICES = ["robot", "screen", "equal"]
TRUST_MIN, TRUST_MAX = 1, 5
ALPHA = 0.05

QUESTIONS_PER_LEVEL = 6          # questions are numbered 1 .. this value (an error if outside)
DESIGN = {                       # design checks: only WARN, the analysis still runs
    "robot_per_level": 3,        # robot advices per level (screen gets the rest)
    "wrong_per_level": 2,        # wrong advices per level
    "wrong_per_source_per_level": 1,   # wrong advices per source per level
    "allow_error_on_q1": False,  # False: warn if question 1 has a wrong advice
    "allow_back_to_back_errors": False,
    "warn_error_on_last_q": True,      # a wrong advice on the last question has no next question (H2, H3)
}
BASELINE_MIN_Q = 3               # H2 baseline: correct advice, not right after an error, question >= this
MIN_CLUSTERS_WARN = 15           # within-person regressions are shaky with fewer participants than this

DATA_COLS = ["participant", "level", "question", "source", "advice", "correct", "chosen", "time_s"]
TRUST_COLS = ["participant", "level", "trust_robot", "trust_screen", "rated_first"]
FINAL_COLS = ["participant", "more_trusted", "reason"]

# Chart colors: robot is always blue, screen always orange, "equal" grey.
COLOR = {"robot": "#2a78d6", "screen": "#eb6834", "equal": "#8a94a0"}
SINGLE = "#2a78d6"
INK, INK_2, GRID = "#1f2328", "#5b6470", "#e6e8eb"


# ==========================================================================
# 1. Reading and checking the Excel file
# ==========================================================================
def read_workbook(path):
    """Reads the three sheets as raw text. Adds excel_row (the row number you see in Excel)."""
    try:
        sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    except FileNotFoundError:
        sys.exit(f"ERROR: file not found: {path}")
    names = {str(n).strip().lower(): n for n in sheets}
    out = {}
    for key, cols in [("data", DATA_COLS), ("trust", TRUST_COLS), ("final", FINAL_COLS)]:
        if key not in names:
            sys.exit(f'ERROR: the workbook has no sheet called "{key.capitalize()}". '
                     f"Sheets found: {list(sheets)}")
        df = sheets[names[key]].copy()
        df.columns = [str(c).strip().lower() for c in df.columns]
        missing = [c for c in cols if c not in df.columns]
        if missing:
            sys.exit(f'ERROR: the {key.capitalize()} sheet is missing these columns: {missing}')
        out[key] = df[cols]
    return out["data"], out["trust"], out["final"]


def txt(s, lower=True):
    """Trim text; empty cells become NaN."""
    def f(v):
        if pd.isna(v):
            return np.nan
        v = str(v).strip()
        if v == "":
            return np.nan
        return v.lower() if lower else v
    return s.map(f).astype(object)


def num(s):
    """Parse numbers (comma decimals allowed). Returns (values, typo_mask)."""
    raw = txt(s).map(lambda v: v.replace(",", ".") if isinstance(v, str) else v)
    val = pd.to_numeric(raw, errors="coerce").astype(float)
    typo = raw.notna() & (val.isna() | ~np.isfinite(val.fillna(0)))
    return val, typo


def _rows(rows, limit=15):
    rows = [int(r) for r in rows]
    shown = ", ".join(str(r) for r in rows[:limit])
    return shown + (f" ... ({len(rows)} rows in total)" if len(rows) > limit else "")


def clean_inputs(data, trust, final):
    """Check and normalise the three raw tables.
    Returns (errors, warnings, data, trust, final). Errors mean the analysis must not run."""
    errors, warns = [], []

    def prep(df, cols):
        df = df.copy()
        if "excel_row" not in df.columns:
            df["excel_row"] = df.index + 2
        empty = df[cols].map(lambda v: pd.isna(v) or str(v).strip() == "").all(axis=1)
        return df[~empty].reset_index(drop=True)

    data, trust, final = prep(data, DATA_COLS), prep(trust, TRUST_COLS), prep(final, FINAL_COLS)

    def bad(sheet, df, mask, msg):
        for r in df.loc[mask, "excel_row"]:
            errors.append(f"{sheet} sheet, row {int(r)}: {msg}")

    # ---- Data ----
    d = data
    for c in ["participant", "level", "source", "advice", "correct", "chosen"]:
        d[c] = txt(d[c])
    d["participant"] = d["participant"].str.upper()
    d["question"], q_typo = num(d["question"])
    d["time_s"], t_typo = num(d["time_s"])
    bad("Data", d, d["participant"].isna(), "participant is empty")
    bad("Data", d, ~d["level"].isin(LEVELS), "level must be easy, medium or hard")
    q_ok = d["question"].between(1, QUESTIONS_PER_LEVEL) & (d["question"] % 1 == 0)
    bad("Data", d, ~q_ok, f"question must be a whole number from 1 to {QUESTIONS_PER_LEVEL}")
    bad("Data", d, ~d["source"].isin(SOURCES), "source must be robot or screen")
    bad("Data", d, ~d["advice"].isin(BOWLS), "advice must be red or blue")
    bad("Data", d, ~d["correct"].isin(BOWLS), "correct must be red or blue")
    bad("Data", d, d["chosen"].notna() & ~d["chosen"].isin(BOWLS), "chosen must be red or blue")
    bad("Data", d, t_typo, "time_s must be a number of seconds, e.g. 4.2 (no units)")
    bad("Data", d, d["time_s"].notna() & (d["time_s"] <= 0), "time_s must be above 0")
    dup = d.duplicated(subset=["participant", "level", "question"], keep=False) & q_ok & d["participant"].notna()
    bad("Data", d, dup, "the same participant, level and question appears more than once")

    # ---- Trust ----
    t = trust
    for c in ["participant", "level", "rated_first"]:
        t[c] = txt(t[c])
    t["participant"] = t["participant"].str.upper()
    bad("Trust", t, t["participant"].isna(), "participant is empty")
    bad("Trust", t, ~t["level"].isin(LEVELS), "level must be easy, medium or hard")
    for c in ["trust_robot", "trust_screen"]:
        v, typo = num(t[c])
        whole_ok = v.between(TRUST_MIN, TRUST_MAX) & (v % 1 == 0)
        bad("Trust", t, typo | (v.notna() & ~whole_ok),
            f"{c} must be a whole number from {TRUST_MIN} to {TRUST_MAX}")
        t[c] = v
    bad("Trust", t, t["rated_first"].notna() & ~t["rated_first"].isin(SOURCES),
        "rated_first must be robot or screen")
    dup = t.duplicated(subset=["participant", "level"], keep=False) & t["level"].isin(LEVELS) & t["participant"].notna()
    bad("Trust", t, dup, "the same participant and level appears more than once")

    # ---- Final ----
    f = final
    for c in ["participant", "more_trusted"]:
        f[c] = txt(f[c])
    f["reason"] = txt(f["reason"], lower=False)
    f["participant"] = f["participant"].str.upper()
    bad("Final", f, f["participant"].isna(), "participant is empty")
    bad("Final", f, ~f["more_trusted"].isin(FINAL_CHOICES), "more_trusted must be robot, screen or equal")
    dup = f.duplicated(subset=["participant"], keep=False) & f["participant"].notna()
    bad("Final", f, dup, "the same participant appears more than once")

    # ---- Cross-sheet: a participant ID in Trust or Final that never appears in Data is almost surely a typo
    ids = set(d["participant"].dropna())
    bad("Trust", t, t["participant"].notna() & ~t["participant"].isin(ids),
        "this participant does not appear in the Data sheet (typo in the ID?)")
    bad("Final", f, f["participant"].notna() & ~f["participant"].isin(ids),
        "this participant does not appear in the Data sheet (typo in the ID?)")

    if errors:
        return errors, warns, data, trust, final

    # ---- Warnings (the analysis still runs) ----
    miss = d.loc[d["chosen"].isna(), "excel_row"]
    if len(miss):
        warns.append(f"Data sheet, row(s) {_rows(miss)}: chosen is empty, left out of the analysis")
    miss = d.loc[d["chosen"].notna() & d["time_s"].isna(), "excel_row"]
    if len(miss):
        warns.append(f"Data sheet, row(s) {_rows(miss)}: time_s is empty, left out of the decision time analysis")
    for c in ["trust_robot", "trust_screen"]:
        miss = t.loc[t[c].isna(), "excel_row"]
        if len(miss):
            warns.append(f"Trust sheet, row(s) {_rows(miss)}: {c} is empty, left out of the analysis")
    miss = t.loc[t["rated_first"].isna(), "excel_row"]
    if len(miss):
        warns.append(f"Trust sheet, row(s) {_rows(miss)}: rated_first is empty, left out of the order check")
    for p in sorted(ids - set(t["participant"])):
        warns.append(f"{p}: has Data rows but no Trust rows")
    for p in sorted(ids - set(f["participant"])):
        warns.append(f"{p}: has Data rows but no row in the Final sheet")
    pairs_data = set(zip(d["participant"], d["level"]))
    for _, r in t.iterrows():
        if (r["participant"], r["level"]) not in pairs_data:
            warns.append(f"Trust sheet, row {int(r['excel_row'])}: {r['participant']} {r['level']} has no Data rows")
    for p, g in d.groupby("participant"):
        missing_levels = set(LEVELS) - set(g["level"])
        if missing_levels:
            warns.append(f"{p}: missing level(s) {sorted(missing_levels)} in the Data sheet")
    warns += design_warnings(d)
    return errors, warns, data, trust, final


def design_warnings(d):
    """Checks that the planned design was really run. Only warnings, and every rule is in DESIGN."""
    w = []
    cfg = DESIGN
    for (p, lvl), g in d.groupby(["participant", "level"]):
        if len(g) != QUESTIONS_PER_LEVEL:
            w.append(f"{p} {lvl}: has {len(g)} questions instead of {QUESTIONS_PER_LEVEL}")
        n_robot = int((g["source"] == "robot").sum())
        exp = cfg["robot_per_level"]
        if exp is not None and n_robot != exp:
            w.append(f"{p} {lvl}: {n_robot} robot and {len(g) - n_robot} screen advices "
                     f"(expected {exp} and {len(g) - exp if len(g) == QUESTIONS_PER_LEVEL else '?'})")
        wrong = g[g["advice"] != g["correct"]]
        if cfg["wrong_per_level"] is not None and len(wrong) != cfg["wrong_per_level"]:
            w.append(f"{p} {lvl}: {len(wrong)} wrong advices (expected {cfg['wrong_per_level']})")
        exp_s = cfg["wrong_per_source_per_level"]
        if exp_s is not None:
            for s in SOURCES:
                n = int((wrong["source"] == s).sum())
                if n != exp_s:
                    w.append(f"{p} {lvl}: {n} wrong advice(s) from the {s} (expected {exp_s})")
        qs = sorted(wrong["question"])
        if not cfg["allow_error_on_q1"] and 1 in qs:
            w.append(f"{p} {lvl}: wrong advice on question 1")
        if not cfg["allow_back_to_back_errors"] and any(b - a == 1 for a, b in zip(qs, qs[1:])):
            w.append(f"{p} {lvl}: two wrong advices back to back")
        if cfg["warn_error_on_last_q"] and QUESTIONS_PER_LEVEL in qs:
            w.append(f"{p} {lvl}: wrong advice on question {QUESTIONS_PER_LEVEL} has no next question, "
                     f"so it cannot be used for H2 and H3")
    return w


# ==========================================================================
# 2. Derived columns
# ==========================================================================
def derive(df):
    df = df.copy()
    df["advice_wrong"] = df["advice"] != df["correct"]
    df["answered"] = df["chosen"].notna()
    df["followed"] = np.where(df["answered"], (df["chosen"] == df["advice"]).astype(float), np.nan)
    df["accurate"] = np.where(df["answered"], (df["chosen"] == df["correct"]).astype(float), np.nan)

    # Previous question in the same level (a new level starts fresh)
    prev = df[["participant", "level", "question", "advice_wrong", "source"]].copy()
    prev["question"] += 1
    prev.columns = ["participant", "level", "question", "prev_wrong", "prev_source"]
    df = df.merge(prev, how="left", on=["participant", "level", "question"])
    df["after_error"] = df["prev_wrong"].fillna(False).astype(bool)
    df["error_source"] = np.where(df["after_error"], df["prev_source"], None)
    df = df.drop(columns=["prev_wrong", "prev_source"])
    # Like with like: only trials whose own advice is correct go into H2 and H3
    df["post_error_trial"] = df["after_error"] & ~df["advice_wrong"]
    df["baseline"] = ~df["advice_wrong"] & ~df["after_error"] & (df["question"] >= BASELINE_MIN_Q)
    df["level"] = pd.Categorical(df["level"], categories=LEVELS, ordered=True)
    return df.sort_values(["participant", "level", "question"]).reset_index(drop=True)


def cell_table(df, trust):
    """One row per participant x level x source: the unit of the trust links."""
    d = df.copy()
    d["followed_wrong"] = d["followed"].where(d["advice_wrong"])
    d["advice_correct"] = (~d["advice_wrong"]).astype(float)
    g = d.groupby(["participant", "level", "source"], observed=True)
    cells = g.agg(n_advice=("question", "size"),
                  follow_rate=("followed", "mean"),
                  fooled=("followed_wrong", "mean"),            # followed that source's wrong advice (1) or caught it (0)
                  own_accuracy=("accurate", "mean"),            # participant's own correct share on this source's questions
                  advice_accuracy=("advice_correct", "mean")).reset_index()
    cells["level"] = cells["level"].astype(str)
    tl = trust.melt(id_vars=["participant", "level", "rated_first"],
                    value_vars=["trust_robot", "trust_screen"], var_name="source", value_name="trust")
    tl["source"] = tl["source"].str.replace("trust_", "", regex=False)
    cells = cells.merge(tl[["participant", "level", "source", "trust", "rated_first"]],
                        how="left", on=["participant", "level", "source"])
    cells["level"] = pd.Categorical(cells["level"], categories=LEVELS, ordered=True)
    return cells


# ==========================================================================
# 3. Statistics helpers
# ==========================================================================
def pct(x):
    return np.nan if pd.isna(x) else round(100 * x, 1)


def mean_ci(x):
    x = pd.Series(x).dropna()
    if len(x) == 0:
        return np.nan, 0.0
    if len(x) < 2:
        return x.mean(), 0.0
    return x.mean(), stats.t.ppf(0.975, len(x) - 1) * stats.sem(x)


def paired_test(a, b, label_a, label_b):
    """Paired t-test (a minus b) with Wilcoxon as backup, on per-participant values."""
    d = pd.concat([pd.Series(a).rename("a"), pd.Series(b).rename("b")], axis=1).dropna()
    n = len(d)
    res = {"n": n, "label_1": label_a, "label_2": label_b,
           "mean_1": d["a"].mean() if n else np.nan, "mean_2": d["b"].mean() if n else np.nan}
    if n < 3:
        res["note"] = "too few participants with data for a test (need at least 3)"
        return res
    diff = d["a"] - d["b"]
    res["estimate"] = diff.mean()
    if diff.std(ddof=1) == 0:
        res["note"] = (f"every participant has exactly the same difference ({diff.mean():+.3f}), "
                       f"so a t-test is not possible")
        return res
    t, p = stats.ttest_rel(d["a"], d["b"])
    ci = stats.t.interval(0.95, n - 1, loc=diff.mean(), scale=stats.sem(diff))
    try:
        w_p = stats.wilcoxon(d["a"], d["b"]).pvalue
    except ValueError:
        w_p = np.nan
    res.update({"ci_low": ci[0], "ci_high": ci[1], "stat": t, "df": n - 1, "p": p, "p_wilcoxon": w_p,
                "effect_size": diff.mean() / diff.std(ddof=1)})
    return res


def welch_test(a, b, label_a, label_b):
    a, b = pd.Series(a).dropna(), pd.Series(b).dropna()
    res = {"n": len(a) + len(b), "label_1": label_a, "label_2": label_b,
           "mean_1": a.mean() if len(a) else np.nan, "mean_2": b.mean() if len(b) else np.nan,
           "n_1": len(a), "n_2": len(b)}
    if len(a) < 3 or len(b) < 3:
        res["note"] = "too few participants in a group for a test (need at least 3 per group)"
        return res
    res["estimate"] = a.mean() - b.mean()
    if a.std(ddof=1) == 0 and b.std(ddof=1) == 0:
        res["note"] = "no variation in either group, so a t-test is not possible"
        return res
    t, p = stats.ttest_ind(a, b, equal_var=False)
    res.update({"stat": t, "p": p})
    return res


def holm(pvals):
    p = np.array(pvals, dtype=float)
    order = np.argsort(p)
    adj = np.empty_like(p)
    running = 0.0
    for rank, i in enumerate(order):
        running = max(running, min(1.0, (len(p) - rank) * p[i]))
        adj[i] = running
    return adj


def within_person_effect(cells, xcol, ycol="trust"):
    """Within-person link between x and trust.
    Linear regression of trust on x with one intercept per participant (so stable differences
    between people drop out) and controls for source and level. Standard errors are clustered by
    participant. Only participants whose x varies across their cells can contribute."""
    res = {"x": xcol}
    d = cells.dropna(subset=[ycol, xcol]).copy()
    d = d[d.groupby("participant")[xcol].transform("size") >= 2]
    d = d[d.groupby("participant")[xcol].transform("nunique") >= 2]
    n_obs, G = len(d), d["participant"].nunique()
    res.update({"n_obs": n_obs, "n": G})
    if G < 3:
        res["note"] = ("too few participants whose values vary across their cells "
                       f"({G}); a test needs at least 3")
        return res
    d["src_robot"] = (d["source"] == "robot").astype(float)
    d["lvl_medium"] = (d["level"].astype(str) == "medium").astype(float)
    d["lvl_hard"] = (d["level"].astype(str) == "hard").astype(float)
    codes = pd.factorize(d["participant"])[0]

    def demean(c):
        return (d[c] - d.groupby("participant")[c].transform("mean")).to_numpy(float)

    x, y = demean(xcol), demean(ycol)
    covs = [c for c in ["src_robot", "lvl_medium", "lvl_hard"] if np.sum(demean(c) ** 2) > 1e-12]
    C = np.column_stack([demean(c) for c in covs]) if covs else np.empty((len(d), 0))
    M = np.column_stack([x, C])
    if np.linalg.matrix_rank(M) < M.shape[1]:
        C, M = np.empty((len(d), 0)), x[:, None]
        covs = []
    if np.linalg.matrix_rank(M) < 1:
        res["note"] = "no usable variation"
        return res
    beta = np.linalg.lstsq(M, y, rcond=None)[0]
    e = y - M @ beta
    bread = np.linalg.pinv(M.T @ M)
    scores = np.zeros((G, M.shape[1]))
    np.add.at(scores, codes, M * e[:, None])
    V = bread @ (scores.T @ scores) @ bread * G / (G - 1)
    se = float(np.sqrt(V[0, 0]))
    res["estimate"] = float(beta[0])
    res["controls"] = "participant" + (", source" if "src_robot" in covs else "") +                       (", level" if any(c.startswith("lvl_") for c in covs) else "")
    # partial correlation (x and trust after removing participant, source and level)
    if C.shape[1]:
        xr = x - C @ np.linalg.lstsq(C, x, rcond=None)[0]
        yr = y - C @ np.linalg.lstsq(C, y, rcond=None)[0]
    else:
        xr, yr = x, y
    res["effect_size"] = float(np.corrcoef(xr, yr)[0, 1]) if xr.std() > 1e-12 and yr.std() > 1e-12 else np.nan
    if not np.isfinite(se) or se <= 1e-12:
        res["note"] = "no residual variation, so the effect cannot be tested"
        return res
    t = beta[0] / se
    tcrit = stats.t.ppf(0.975, G - 1)
    res.update({"ci_low": beta[0] - tcrit * se, "ci_high": beta[0] + tcrit * se, "stat": t, "df": G - 1,
                "p": 2 * stats.t.sf(abs(t), G - 1)})
    if G < MIN_CLUSTERS_WARN:
        res["caution"] = f"only {G} participants contribute, standard errors may be too small"
    return res


def permutation_post_vs_base(df, n_perm=4000, seed=1):
    """H2 p-value. Statistic: mean over participants of (follow rate after an error minus baseline follow rate).
    Under 'no effect' the labels after_error / baseline are exchangeable within a participant and level,
    so they are shuffled there (counts per participant and level stay the same)."""
    d = df[(df["post_error_trial"] | df["baseline"]) & df["answered"]]
    rng = np.random.default_rng(seed)
    obs, perm = [], np.zeros(n_perm)
    groups = [g for _, g in d.groupby("participant") if g["post_error_trial"].any() and g["baseline"].any()]
    if len(groups) < 3:
        return np.nan
    for g in groups:
        y = g["followed"].to_numpy(float)
        lab = g["post_error_trial"].to_numpy(bool)
        L = np.empty((n_perm, len(g)), dtype=bool)
        lv = g["level"].astype(str).to_numpy()
        for s in np.unique(lv):
            idx = np.where(lv == s)[0]
            L[:, idx] = rng.permuted(np.tile(lab[idx], (n_perm, 1)), axis=1)
        n_post = lab.sum()
        obs.append(y[lab].mean() - y[~lab].mean())
        perm += (L @ y) / n_post - ((~L) @ y) / (len(y) - n_post)
    stat, perm = np.mean(obs), perm / len(groups)
    return (1 + np.sum(np.abs(perm) >= abs(stat) - 1e-12)) / (n_perm + 1)


def per_participant(df, value, by, mask=None):
    d = df if mask is None else df[mask]
    d = d[d[value].notna()]
    return d.groupby(["participant"] + by, observed=True)[value].mean().unstack(by)


# ==========================================================================
# 4. All analyses
# ==========================================================================
def compute(df, trust, final):
    """df: derived Data. Returns dict R with tables (for Excel), tests (list of dicts) and extras."""
    R = {"T": {}, "tests": [], "x": {}}
    T, tests = R["T"], R["tests"]
    a = df[df["answered"]]
    cells = cell_table(df, trust)
    R["cells"], R["df"], R["trust"], R["final"] = cells, df, trust, final

    def add(tid, role, question, measure, test, res):
        res = dict(res)
        res.update({"id": tid, "role": role, "question": question, "measure": measure, "test": test})
        p = res.get("p", np.nan)
        res["significant"] = bool(pd.notna(p) and p < ALPHA)
        tests.append(res)
        return res

    # ---- Per-participant values ----
    tr = trust.melt(id_vars=["participant", "level"], value_vars=["trust_robot", "trust_screen"],
                    var_name="source", value_name="trust")
    tr["source"] = tr["source"].str.replace("trust_", "", regex=False)
    tr["level"] = pd.Categorical(tr["level"], categories=LEVELS, ordered=True)
    trust_pp = tr.dropna(subset=["trust"]).groupby(["participant", "source"])["trust"].mean().unstack()
    trust_pp = trust_pp.reindex(columns=SOURCES)
    fol = per_participant(df, "followed", ["source"]).reindex(columns=SOURCES)
    R["x"].update({"n_participants": df["participant"].nunique(), "n_answered": len(a), "n_rows": len(df),
                   "trust_pp": trust_pp, "fol_pp": fol})

    # ---- P1: which source is trusted more ----
    r = paired_test(trust_pp["robot"], trust_pp["screen"], "robot", "screen")
    add("P1", "Primary", "Which source is trusted more?",
        "Mean trust rating over the 3 levels (1 to 5), robot minus screen", "Paired t-test (Wilcoxon backup)", r)

    # ---- P2: follow rate robot vs screen (H1 behavior) ----
    r = paired_test(fol["robot"], fol["screen"], "robot", "screen")
    add("P2", "Primary", "Is the robot followed more than the screen? (H1, behavior)",
        "Share of advices followed, robot minus screen", "Paired t-test (Wilcoxon backup)", r)

    # ---- P3: trust vs following (within person) ----
    r = within_person_effect(cells, "follow_rate")
    add("P3", "Primary", "Is a source followed more when it is trusted more?",
        "Trust (1 to 5) per level and source vs share of that source's advices followed (0 to 1)",
        "Within-person regression, cluster-robust", r)

    # ---- S1: final overall choice ----
    fin = final.copy()
    counts = fin["more_trusted"].value_counts().reindex(FINAL_CHOICES).fillna(0).astype(int)
    k, n_dec = int(counts["robot"]), int(counts["robot"] + counts["screen"])
    res = {"n": int(counts.sum()), "label_1": "robot", "label_2": "screen", "n_robot": k,
           "n_screen": int(counts["screen"]), "n_equal": int(counts["equal"])}
    if n_dec >= 1:
        res["estimate"] = k / n_dec
        res["ci_low"], res["ci_high"] = stats.binomtest(k, n_dec, 0.5).proportion_ci(method="exact")
        res["p"] = stats.binomtest(k, n_dec, 0.5).pvalue
        res["stat"] = k
    else:
        res["note"] = "nobody chose robot or screen, so no test is possible"
    add("S1", "Secondary", "Which source was trusted more overall (final question)?",
        "Share choosing the robot among those who chose robot or screen (equal answers left out)",
        "Exact binomial test against 50%", res)
    T["Final choice"] = pd.DataFrame({"more_trusted": FINAL_CHOICES, "participants": counts.values,
                                      "percent": [pct(c / counts.sum()) if counts.sum() else np.nan for c in counts]})
    diff_pp = trust_pp["robot"] - trust_pp["screen"]
    rating_says = diff_pp.map(lambda v: np.nan if pd.isna(v) else ("robot" if v > 1e-9 else "screen" if v < -1e-9 else "equal"))
    x = fin.set_index("participant")["more_trusted"].to_frame().join(rating_says.rename("ratings_say"), how="inner")
    x = x.dropna()
    R["x"]["agree"] = (int((x["more_trusted"] == x["ratings_say"]).sum()), len(x))
    T["Final vs ratings"] = pd.crosstab(x["more_trusted"], x["ratings_say"]).reindex(
        index=FINAL_CHOICES, columns=FINAL_CHOICES, fill_value=0).rename_axis(index="final answer", columns="ratings say")\
        .reset_index()
    T["Final reasons"] = fin[["participant", "more_trusted", "reason"]].sort_values("participant")

    # ---- S2, S3: trust and correctness ----
    r2 = within_person_effect(cells, "fooled")
    r3 = within_person_effect(cells, "own_accuracy")
    ps = [r.get("p", np.nan) for r in (r2, r3)]
    ok = [i for i, p in enumerate(ps) if pd.notna(p)]
    if ok:
        adj = holm([ps[i] for i in ok])
        for i, a_ in zip(ok, adj):
            (r2, r3)[i]["p_holm"] = a_
    r2["significant_holm"] = bool(pd.notna(r2.get("p_holm")) and r2["p_holm"] < ALPHA)
    r3["significant_holm"] = bool(pd.notna(r3.get("p_holm")) and r3["p_holm"] < ALPHA)
    add("S2", "Secondary", "Does trust in a source drop when the participant followed its wrong advice?",
        "Trust per level and source vs followed that source's wrong advice (1) or not (0)",
        "Within-person regression, Holm over S2 and S3", r2)
    add("S3", "Secondary", "Does trust follow the participant's own right and wrong answers?",
        "Trust per level and source vs the participant's own share correct on that source's questions",
        "Within-person regression, Holm over S2 and S3", r3)
    adv = cells["advice_accuracy"].dropna()
    R["x"]["advice_accuracy_varies"] = bool(len(adv) and adv.nunique() > 1)
    R["x"]["advice_accuracy_values"] = sorted(adv.round(3).unique().tolist())
    if R["x"]["advice_accuracy_varies"]:
        add("S3b", "Secondary", "Does trust follow the correctness of the source's advice?",
            "Trust per level and source vs share of that source's advices that were correct",
            "Within-person regression", within_person_effect(cells, "advice_accuracy"))

    # ---- S4: between-person link ----
    fdiff = fol["robot"] - fol["screen"]
    dd = pd.concat([diff_pp.rename("t"), fdiff.rename("f")], axis=1).dropna()
    res = {"n": len(dd)}
    if len(dd) < 4 or dd["t"].nunique() < 2 or dd["f"].nunique() < 2:
        res["note"] = "too few participants or no variation, so no correlation is possible"
    else:
        rho, p = stats.spearmanr(dd["t"], dd["f"])
        res.update({"estimate": rho, "p": p, "effect_size": rho})
    add("S4", "Secondary", "Do people who trust the robot more than the screen also follow it more?",
        "Trust difference (robot minus screen) vs follow-rate difference, one point per person",
        "Spearman correlation", res)
    R["x"]["diff_pp"] = dd

    # ---- S5: H2 ----
    post = df[df["post_error_trial"] & df["answered"]].groupby("participant")["followed"].mean()
    base = df[df["baseline"] & df["answered"]].groupby("participant")["followed"].mean()
    res = paired_test(post, base, "after_error", "baseline")
    if "p" in res:      # near-ceiling, lumpy proportions: the t-test is too liberal here, so p comes from permutations
        res["p_t_test"] = res["p"]
        res["p"] = permutation_post_vs_base(df)
    add("S5", "Secondary", "H2: Is advice followed less right after a wrong advice?",
        f"Share followed on correct advice right after a wrong advice vs baseline (question {BASELINE_MIN_Q}+, not after an error)",
        "Within-person permutation test (labels shuffled within participant and level); t-based CI", res)
    R["x"].update({"post": post, "base": base})

    # ---- S6: H3 ----
    post_src = per_participant(df, "followed", ["error_source"], df["post_error_trial"] & df["answered"])
    post_src = post_src.reindex(columns=SOURCES)
    add("S6", "Exploratory", "H3: Is the drop different after robot errors than after screen errors?",
        "Share followed right after a robot error vs right after a screen error",
        "Paired t-test (Wilcoxon backup)", paired_test(post_src["robot"], post_src["screen"], "after_robot_error",
                                                       "after_screen_error"))
    R["x"]["post_src"] = post_src

    # ---- S7: H4 ----
    fol_lvl = per_participant(df, "followed", ["level"]).reindex(columns=LEVELS)
    res = paired_test(fol_lvl["hard"], fol_lvl["easy"], "hard", "easy")
    w = fol_lvl.dropna()
    res["mean_medium"] = w["medium"].mean() if len(w) else np.nan
    try:
        res["p_friedman"] = stats.friedmanchisquare(*[w[l] for l in LEVELS]).pvalue if len(w) >= 3 else np.nan
    except ValueError:
        res["p_friedman"] = np.nan
    add("S7", "Secondary", "H4: Is advice followed more as the levels get harder?",
        "Share followed, hard minus easy (medium shown in the tables)", "Paired t-test (Friedman backup)", res)
    R["x"]["fol_lvl"] = fol_lvl

    # ---- C1: difficulty checks ----
    acc_lvl = per_participant(df, "accurate", ["level"]).reindex(columns=LEVELS)
    add("C1a", "Check", "Difficulty check: is accuracy lower in the hard level?",
        "Share of correct bowls, hard minus easy", "Paired t-test (Wilcoxon backup)",
        paired_test(acc_lvl["hard"], acc_lvl["easy"], "hard", "easy"))
    d_time = df[df["answered"] & df["time_s"].notna()].copy()
    d_time["log_time"] = np.log(d_time["time_s"])
    time_lvl = per_participant(d_time, "log_time", ["level"]).reindex(columns=LEVELS)
    res = paired_test(time_lvl["hard"], time_lvl["easy"], "hard", "easy")
    for k_ in ("mean_1", "mean_2"):
        res[k_] = np.exp(res[k_]) if pd.notna(res.get(k_)) else np.nan      # geometric mean in seconds
    res["time_note"] = "means are geometric means in seconds; the test is on log time (estimate and CI are log differences)"
    add("C1b", "Check", "Difficulty check: are decisions slower in the hard level?",
        "Decision time (log seconds), hard minus easy", "Paired t-test (Wilcoxon backup)", res)
    R["x"].update({"acc_lvl": acc_lvl, "time_lvl": time_lvl})

    # ---- C2: rating order ----
    rf = trust.dropna(subset=["rated_first"]).copy()
    rf["diff"] = rf["trust_robot"] - rf["trust_screen"]
    rf = rf.dropna(subset=["diff"])
    per = rf.groupby(["participant", "rated_first"])["diff"].mean().unstack()
    for s in SOURCES:
        if s not in per:
            per[s] = np.nan
    both = per.dropna(subset=SOURCES)
    if len(both) >= 3:
        res = paired_test(per.loc[both.index, "robot"], per.loc[both.index, "screen"], "robot_rated_first",
                          "screen_rated_first")
        res["design"] = "within participants (rated_first changes between levels)"
    else:
        res = welch_test(per["robot"], per["screen"], "robot_rated_first", "screen_rated_first")
        res["design"] = "between participants (rated_first fixed per person)"
    res["n_robot_first"] = int(per["robot"].notna().sum())
    res["n_screen_first"] = int(per["screen"].notna().sum())
    add("C2", "Check", "Order check: does the trust difference depend on which source was rated first?",
        "Trust difference (robot minus screen) when robot was rated first vs screen first",
        "Welch t-test (or paired if the order changes within people)", res)

    # ---- Descriptive tables ----
    tl = tr.dropna(subset=["trust"])
    T["Trust by source"] = tl.groupby("source")["trust"].agg(mean="mean", sd="std", n_ratings="size").reindex(SOURCES).reset_index()
    tb = tl.groupby(["level", "source"], observed=True)["trust"].agg(mean="mean", sd="std", n_ratings="size").reset_index()
    T["Trust by level"] = tb
    wide = tb.pivot(index="level", columns="source", values="mean").reindex(LEVELS)
    T["Trust by level"] = tb.merge(
        (wide["robot"] - wide["screen"]).rename("robot_minus_screen").reset_index(), on="level", how="left")

    def summarize(g):
        gw, gc = g[g["advice_wrong"]], g[~g["advice_wrong"]]
        return pd.Series({
            "questions": len(g),
            "followed_%": pct(g["followed"].mean()),
            "accuracy_%": pct(g["accurate"].mean()),
            "followed_wrong_advice_%": pct(gw["followed"].mean()),
            "ignored_correct_advice_%": pct(1 - gc["followed"].mean()) if len(gc) else np.nan,
            "mean_time_s": round(g["time_s"].mean(), 2),
            "median_time_s": round(g["time_s"].median(), 2)})

    def table(by):
        return pd.DataFrame([dict(zip(by if isinstance(by, list) else [by], k if isinstance(k, tuple) else (k,)),
                                  **summarize(g).to_dict())
                             for k, g in a.groupby(by, observed=True)])
    T["Reliance by level"] = table("level")
    T["Reliance by source"] = table("source")
    T["Reliance level x source"] = table(["level", "source"])

    tally = a.pivot_table(index="participant", columns="level", values="accurate", aggfunc="sum",
                          observed=True).reindex(columns=LEVELS)
    tally.columns = [f"{c}_correct_of_{QUESTIONS_PER_LEVEL}" for c in tally.columns]
    pp = tally.join(trust_pp.add_prefix("mean_trust_"))
    pp["followed_robot_%"] = (fol["robot"] * 100).round(1)
    pp["followed_screen_%"] = (fol["screen"] * 100).round(1)
    pp["final_choice"] = fin.set_index("participant")["more_trusted"]
    pp["mean_time_s"] = a.groupby("participant")["time_s"].mean().round(2)
    T["Per participant"] = pp.reset_index()

    ht = pd.DataFrame(tests)
    first = ["id", "role", "question", "measure", "test", "n", "label_1", "mean_1", "label_2", "mean_2", "estimate",
             "ci_low", "ci_high", "stat", "df", "p", "p_holm", "p_t_test", "p_wilcoxon", "p_friedman", "effect_size", "significant"]
    ht = ht[[c for c in first if c in ht] + [c for c in ht if c not in first]]
    T = {"Tests": ht, **T}
    T["Link data (cells)"] = cells.assign(level=cells["level"].astype(str))
    T["Cleaned data"] = df.drop(columns=["excel_row"], errors="ignore").assign(level=df["level"].astype(str))
    R["T"] = T
    return R


# ==========================================================================
# 5. Plain-language summary
# ==========================================================================
def fmt_p(p):
    if pd.isna(p):
        return "p = n/a"
    return "p < .001" if p < 0.001 else f"p = {p:.3f}"


def find(R, tid):
    return next((t for t in R["tests"] if t["id"] == tid), None)


def paired_lines(r, as_pct=False, digits=2, unit=""):
    f = (lambda v: f"{100 * v:.1f}%") if as_pct else (lambda v: f"{v:.{digits}f}{unit}")
    out = []
    if pd.notna(r.get("mean_1")):
        out.append(f"   {r['label_1']}: {f(r['mean_1'])}   {r['label_2']}: {f(r['mean_2'])}   (n = {r['n']} participants)")
    if "note" in r:
        out.append(f"   Note: {r['note']}")
    elif "stat" in r:
        dif = (f"{100 * r['estimate']:+.1f} percentage points" if as_pct else f"{r['estimate']:+.{digits}f}")
        ci = (f"[{100 * r['ci_low']:+.1f}, {100 * r['ci_high']:+.1f}]" if as_pct
              else f"[{r['ci_low']:+.{digits}f}, {r['ci_high']:+.{digits}f}]")
        out.append(f"   difference {dif}, 95% CI {ci}")
        if "p_t_test" in r:
            out.append(f"   within-person permutation test {fmt_p(r['p'])}, dz = {r['effect_size']:.2f}"
                       f"   (paired t-test {fmt_p(r['p_t_test'])}, shown for comparison only)")
        else:
            out.append(f"   paired t({r['df']}) = {r['stat']:.2f}, {fmt_p(r['p'])}, dz = {r['effect_size']:.2f}"
                       f"   (Wilcoxon {fmt_p(r.get('p_wilcoxon'))})")
    return out


def verdict_paired(r, what, name_a, name_b, expected=None):
    """expected: None, 'a_higher' or 'a_lower' (direction predicted by the hypothesis)."""
    if "stat" not in r:
        return "   Result: no test possible with these data."
    if not r["significant"]:
        return (f"   Result: no reliable difference in {what} between {name_a} and {name_b} "
                f"(this does not prove they are equal).")
    a_higher = r["estimate"] > 0
    s = f"   Result: {what} was reliably {'higher' if a_higher else 'lower'} for {name_a} than for {name_b}"
    if expected:
        s += " (as predicted)." if (expected == "a_higher") == a_higher else " (OPPOSITE to the prediction)."
    else:
        s += "."
    return s


def link_lines(r, xname, unit_phrase, direction_pos, direction_neg):
    out = [f"   participants contributing: {r.get('n', 0)} ({r.get('n_obs', 0)} level-and-source cells)"]
    if "note" in r:
        out.append(f"   Note: {r['note']}")
        return out
    out.append(f"   change in trust (1 to 5 scale) {unit_phrase}: {r['estimate']:+.2f}, "
               f"95% CI [{r['ci_low']:+.2f}, {r['ci_high']:+.2f}], t({r['df']}) = {r['stat']:.2f}, {fmt_p(r['p'])}"
               + (f", Holm-adjusted {fmt_p(r['p_holm'])}" if "p_holm" in r and pd.notna(r["p_holm"]) else "")
               + f", partial r = {r['effect_size']:.2f}")
    sig = r["significant_holm"] if "significant_holm" in r else r["significant"]
    if not sig:
        out.append(f"   Result: no reliable link between trust and {xname}.")
    else:
        out.append(f"   Result: {direction_pos if r['estimate'] > 0 else direction_neg}")
    if "caution" in r:
        out.append(f"   Caution: {r['caution']}.")
    return out


def build_summary(R, warns):
    x, S = R["x"], []
    P1, P2, P3, S1, S2, S3, S4, S5, S6, S7 = [find(R, k) for k in
                                              ["P1", "P2", "P3", "S1", "S2", "S3", "S4", "S5", "S6", "S7"]]
    C1a, C1b, C2 = find(R, "C1a"), find(R, "C1b"), find(R, "C2")
    S += ["HRI ROBOT TRUST STUDY: RESULTS SUMMARY", "=" * 40, "",
          f"Participants: {x['n_participants']}    Questions answered: {x['n_answered']} of {x['n_rows']}",
          "A result counts as significant when p < .05. 'Reliable' below means significant.",
          "Trust is rated 1 (lowest) to 5 (highest). Following = the chosen bowl matched the advised bowl.", ""]

    S += ["PART 1: TRUST (main question)", "-" * 40, ""]
    S += ["1. Which source do participants trust more? (primary test P1)"]
    S += paired_lines(P1, digits=2)
    S += [verdict_paired(P1, "trust", "the robot", "the screen", "a_higher"), ""]

    S += ["2. Which source did participants say they trusted more overall? (final question, S1)"]
    S += [f"   robot: {S1['n_robot']}   screen: {S1['n_screen']}   equal: {S1['n_equal']}"]
    if "note" in S1:
        S += [f"   Note: {S1['note']}"]
    else:
        S += [f"   Among the {S1['n_robot'] + S1['n_screen']} who chose a source, {100 * S1['estimate']:.0f}% chose the robot "
              f"(95% CI {100 * S1['ci_low']:.0f}% to {100 * S1['ci_high']:.0f}%), exact binomial {fmt_p(S1['p'])}."]
        if not S1["significant"]:
            S += ["   Result: no reliable preference for either source."]
        else:
            S += [f"   Result: reliably more participants chose the {'robot' if S1['estimate'] > 0.5 else 'screen'}."]
    ag, tot = x["agree"]
    S += [f"   The final answer matched the average ratings for {ag} of {tot} participants. "
          f"The reasons are listed in the results Excel (sheet 'Final reasons')."]
    S += [""]

    S += ["3. Trust by level (descriptive, no test)"]
    tb = R["T"]["Trust by level"]
    for lvl in LEVELS:
        g = tb[tb["level"] == lvl].set_index("source")["mean"]
        if len(g) == 2:
            S += [f"   {lvl}: robot {g['robot']:.2f}   screen {g['screen']:.2f}   (robot minus screen {g['robot'] - g['screen']:+.2f})"]
    S += [""]

    S += ["4. Order check (C2): did the rating order matter?"]
    S += [f"   participants who rated robot first: {C2['n_robot_first']}, screen first: {C2['n_screen_first']}"]
    if "stat" in C2:
        S += [f"   trust difference (robot minus screen): robot first {C2['mean_1']:+.2f}, screen first {C2['mean_2']:+.2f}, "
              f"{fmt_p(C2['p'])}"]
        S += ["   Result: " + ("the rating order may have affected the ratings, treat the trust comparison with care."
                               if C2["significant"] else "no sign that the rating order changed the trust difference.")]
    else:
        S += [f"   Note: {C2.get('note', 'no test possible')}"]
    S += [""]

    S += ["PART 2: TRUST AND FOLLOWING THE ADVICE", "-" * 40, ""]
    S += ["5. Is the robot followed more than the screen? (primary test P2, H1 behavior)"]
    S += paired_lines(P2, as_pct=True)
    S += [verdict_paired(P2, "the share of advice followed", "the robot", "the screen", "a_higher"), ""]
    S += ["6. When a source is trusted more, is it also followed more? (primary test P3)",
          "   Compares the same person across their level-and-source cells, accounting for source and level."]
    S += link_lines(P3, "following", "for each extra 100% of that source's advice followed",
                    "trust is higher when the source was followed more.", "trust is lower when the source was followed more.")
    S += ["   This shows an association only. Trust is rated after the level, so following came first."]
    S += [""]
    S += ["7. Between people (S4): do those who trust the robot more than the screen also follow it more?"]
    if "estimate" in S4:
        S += [f"   Spearman rho = {S4['estimate']:.2f} (n = {S4['n']}), {fmt_p(S4['p'])}"]
        S += ["   Result: " + ("reliable link between the two differences." if S4["significant"] else "no reliable link.")]
    else:
        S += [f"   Note: {S4['note']}"]
    S += [""]

    S += ["PART 3: TRUST AND CORRECTNESS", "-" * 40, ""]
    vals = x["advice_accuracy_values"]
    if not x["advice_accuracy_varies"]:
        S += [f"Advice correctness: each source's advice was correct {', '.join(f'{100 * v:.0f}%' for v in vals)} of the time "
              f"in every level, by design, so it cannot explain trust differences."]
    else:
        S += ["Advice correctness varied across level-and-source cells, see test S3b in the Tests sheet."]
        s3b = find(R, "S3b")
        S += link_lines(s3b, "the correctness of the advice", "for each extra 100% of the source's advice being correct",
                        "trust is higher when the source's advice was more often correct.",
                        "trust is lower when the source's advice was more often correct.")
    S += [""]
    S += ["8. Does trust drop when the participant followed that source's wrong advice? (S2)"]
    S += link_lines(S2, "being fooled by the source", "when the wrong advice was followed (vs caught)",
                    "trust was higher after following the wrong advice.", "trust was lower after following the wrong advice.")
    S += [""]
    S += ["9. Does trust follow the participant's own correctness with that source's questions? (S3)"]
    S += link_lines(S3, "own correctness", "for each extra 100% correct on that source's questions",
                    "trust was higher when the participant was more often right.",
                    "trust was lower when the participant was more often right.")
    S += ["   S2 and S3 are adjusted together (Holm), because they are exploratory.", ""]

    S += ["PART 4: SECONDARY BEHAVIOR HYPOTHESES (H2 to H4) AND CHECKS", "-" * 40, ""]
    S += ["H2 (S5): Is advice followed less right after a wrong advice?"]
    S += paired_lines(S5, as_pct=True)
    S += [verdict_paired(S5, "the share followed", "right after an error", "baseline", "a_lower"), ""]
    S += ["H3 (S6, exploratory): Is the drop different after robot errors than after screen errors?"]
    S += paired_lines(S6, as_pct=True)
    S += [verdict_paired(S6, "the share followed", "after robot errors", "after screen errors"),
          "   (Based on very few trials per person, so low power.)", ""]
    S += ["H4 (S7): Is advice followed more as the levels get harder?"]
    fl = x["fol_lvl"].mean()
    S += [f"   easy {100 * fl['easy']:.1f}%   medium {100 * fl['medium']:.1f}%   hard {100 * fl['hard']:.1f}%"]
    S += paired_lines(S7, as_pct=True)[1:]
    if "p_friedman" in S7 and pd.notna(S7["p_friedman"]):
        S += [f"   Friedman test across all three levels: {fmt_p(S7['p_friedman'])}"]
    S += [verdict_paired(S7, "the share followed", "the hard level", "the easy level", "a_higher"), ""]
    S += ["Difficulty checks (are the levels really harder?)"]
    al = x["acc_lvl"].mean()
    S += [f"   accuracy: easy {100 * al['easy']:.1f}%   medium {100 * al['medium']:.1f}%   hard {100 * al['hard']:.1f}%"]
    S += paired_lines(C1a, as_pct=True)[1:]
    S += [verdict_paired(C1a, "accuracy", "the hard level", "the easy level", "a_lower")]
    tl_ = np.exp(x["time_lvl"].mean())
    if tl_.notna().all():
        S += [f"   typical decision time: easy {tl_['easy']:.1f}s   medium {tl_['medium']:.1f}s   hard {tl_['hard']:.1f}s"]
        if "stat" in C1b:
            S += [f"   hard decisions took {np.exp(C1b['estimate']):.2f} times as long as easy ones "
                  f"(95% CI {np.exp(C1b['ci_low']):.2f} to {np.exp(C1b['ci_high']):.2f}), paired t({C1b['df']}) = "
                  f"{C1b['stat']:.2f}, {fmt_p(C1b['p'])}   (Wilcoxon {fmt_p(C1b.get('p_wilcoxon'))})"]
        elif "note" in C1b:
            S += [f"   Note: {C1b['note']}"]
        S += [verdict_paired(C1b, "decision time", "the hard level", "the easy level", "a_higher")]
    else:
        S += ["   decision time: no times entered, so this check was skipped."]
    S += [""]
    S += ["Notes:",
          "- Paired tests compare each person with themselves. A Wilcoxon test is reported as a backup.",
          "- H2 uses a permutation test: follow rates of correct advice sit near 100%, where the t-test is too liberal.",
          "- Following correct advice can also mean the participant knew the answer, so 'followed' mixes",
          "  reliance with agreement. Following WRONG advice (Reliance sheets) is the cleaner reliance signal.",
          "- Levels always run easy, medium, hard, so difficulty is partly mixed up with fatigue and learning.",
          "- No correction for multiple tests on the three primary tests. S2 and S3 are adjusted with Holm.",
          "- H3 is exploratory (few trials per person). See the README for limitations and assumptions."]
    if warns:
        S += ["", f"There were {len(warns)} data warnings. They are printed on screen and listed in the Excel file."]
    return "\n".join(S)


# ==========================================================================
# 6. Charts
# ==========================================================================
def charts(R, outdir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 10, "axes.edgecolor": GRID, "axes.labelcolor": INK_2,
        "xtick.color": INK_2, "ytick.color": INK_2, "axes.spines.top": False, "axes.spines.right": False,
        "axes.grid": True, "axes.grid.axis": "y", "grid.color": GRID, "grid.linewidth": 0.8,
        "axes.axisbelow": True, "axes.titlesize": 12, "axes.titleweight": "bold", "axes.titlecolor": INK,
        "axes.titlelocation": "left", "legend.frameon": False})
    x, df, cells = R["x"], R["df"], R["cells"]
    files, rng = [], np.random.default_rng(1)

    def save(fig, name):
        f = os.path.join(outdir, name)
        fig.savefig(f, dpi=160)
        plt.close(fig)
        files.append(f)

    def cat_points(g, col, values, off):
        """Mean trust per category. A mean needs >= 3 cells, its interval >= 8 (small groups give silly intervals)."""
        out = []
        for v in values:
            t_ = g[g[col].round(2) == v]["trust"]
            if len(t_) >= 3:
                out.append((v + off, t_.mean(), mean_ci(t_)[1] if len(t_) >= 8 else 0))
        return out

    def legend_handles(keys):
        return [Line2D([0], [0], marker="o", color="none", markerfacecolor=COLOR[k], markersize=8,
                       label=k.capitalize()) for k in keys]

    def note(fig, text):
        fig.text(0.01, 0.01, text, color=INK_2, fontsize=8)

    # 1: trust robot vs screen per participant
    tp = x["trust_pp"].dropna()
    if len(tp):
        fig, ax = plt.subplots(figsize=(5.6, 4.6))
        for _, r_ in tp.iterrows():
            ax.plot([0, 1], [r_["screen"], r_["robot"]], color="#c9ced4", lw=0.9, zorder=1)
        for i, s in enumerate(["screen", "robot"]):
            m, ci = mean_ci(tp[s])
            ax.errorbar(i, m, yerr=ci, fmt="o", color=COLOR[s], ecolor=COLOR[s], markersize=11, capsize=5,
                        elinewidth=2.5, markeredgecolor="white", markeredgewidth=1.5, zorder=3)
            ax.text(i + (0.07 if i else -0.07), m, f"{m:.2f}", color=INK, fontsize=11, fontweight="bold",
                    ha="left" if i else "right", va="center")
        ax.set_xticks([0, 1], ["Screen", "Robot"]); ax.set_xlim(-0.4, 1.4)
        ax.set_ylim(TRUST_MIN - 0.2, TRUST_MAX + 0.2); ax.set_ylabel("Trust (mean of 3 levels, 1 to 5)")
        ax.set_title("Trust in the robot vs the screen")
        ax.grid(False); ax.grid(True, axis="y")
        note(fig, "Grey lines: one participant each. Dots: group mean with 95% confidence interval.")
        fig.tight_layout(rect=(0, 0.03, 1, 1)); save(fig, "1_trust_robot_vs_screen.png")

    # 2: trust by level
    tl = R["trust"].melt(id_vars=["participant", "level"], value_vars=["trust_robot", "trust_screen"],
                         var_name="source", value_name="trust")
    tl["source"] = tl["source"].str.replace("trust_", "", regex=False)
    if tl["trust"].notna().any():
        fig, ax = plt.subplots(figsize=(6.4, 4.4))
        xs = np.arange(3)
        for s, off in zip(SOURCES, (-0.03, 0.03)):
            vals = [mean_ci(tl[(tl["level"] == l) & (tl["source"] == s)]["trust"]) for l in LEVELS]
            m = np.array([v[0] for v in vals]); c = np.array([v[1] for v in vals])
            ax.errorbar(xs + off, m, yerr=c, color=COLOR[s], marker="o", markersize=8, lw=2, capsize=4,
                        markeredgecolor="white", markeredgewidth=1.5, label=s.capitalize())
        ax.set_xticks(xs, [l.capitalize() for l in LEVELS]); ax.set_xlim(-0.3, 2.3)
        ax.set_ylim(TRUST_MIN - 0.2, TRUST_MAX + 0.2); ax.set_ylabel("Trust (1 to 5)")
        ax.set_title("Trust after each level"); ax.legend(handles=legend_handles(SOURCES), loc="lower left", ncol=2)
        note(fig, "Points: mean across participants, lines: 95% confidence interval.")
        fig.tight_layout(rect=(0, 0.03, 1, 1)); save(fig, "2_trust_by_level.png")

    # 3: final choice
    fc = R["T"]["Final choice"]
    if fc["participants"].sum() > 0:
        fig, ax = plt.subplots(figsize=(6, 3.6))
        ypos = np.arange(3)[::-1]
        ax.barh(ypos, fc["participants"], 0.55, color=[COLOR[k] for k in FINAL_CHOICES])
        for y, v in zip(ypos, fc["participants"]):
            ax.text(v + 0.15, y, str(int(v)), va="center", color=INK, fontweight="bold")
        ax.set_yticks(ypos, [k.capitalize() for k in FINAL_CHOICES]); ax.grid(False); ax.grid(True, axis="x")
        ax.set_xlabel("Participants"); ax.set_title("Who was trusted more overall?")
        ax.set_xlim(0, fc["participants"].max() * 1.15 + 1)
        fig.tight_layout(); save(fig, "3_final_choice.png")

    # 4: trust vs following
    c4 = cells.dropna(subset=["trust", "follow_rate"]).copy()
    if len(c4):
        fig, ax = plt.subplots(figsize=(6.8, 4.6))
        c4["fr"] = c4["follow_rate"].round(2)
        levels_x = sorted(c4["fr"].unique())
        for s, off in zip(SOURCES, (-0.06, 0.06)):
            g = c4[c4["source"] == s]
            ax.scatter(g["fr"] + off + rng.uniform(-0.025, 0.025, len(g)), g["trust"] + rng.uniform(-0.12, 0.12, len(g)),
                       s=9, color=COLOR[s], alpha=0.22, linewidths=0, zorder=1)
            pts = cat_points(c4.assign(fr=c4["fr"])[c4["source"] == s], "fr", levels_x, off)
            ax.errorbar([p[0] for p in pts], [p[1] for p in pts], yerr=[p[2] for p in pts],
                        color=COLOR[s], marker="o", markersize=8, lw=2, capsize=3, markeredgecolor="white",
                        markeredgewidth=1.5, zorder=3, label=s.capitalize())
        ax.set_xticks(levels_x, [f"{100 * v:.0f}%" for v in levels_x])
        ax.set_xlabel("Share of that source's advices followed in the level")
        ax.set_ylabel("Trust (1 to 5)"); ax.set_ylim(TRUST_MIN - 0.4, TRUST_MAX + 0.4)
        ax.set_title("Trust vs following the advice"); ax.legend(handles=legend_handles(SOURCES), loc="lower right", ncol=2)
        note(fig, "Faint dots: one level and source of one participant. Bold: mean (shown when 3+ cells), 95% CI when 8+ cells (rough guide).")
        fig.tight_layout(rect=(0, 0.03, 1, 1)); save(fig, "4_trust_vs_following.png")

    # 5: trust vs correctness (two panels, same scale)
    c5 = cells.dropna(subset=["trust"])
    if len(c5):
        fig, axes = plt.subplots(1, 2, figsize=(10, 4.4), sharey=True)
        ax = axes[0]
        f = c5.dropna(subset=["fooled"])
        for s, off in zip(SOURCES, (-0.08, 0.08)):
            pts = cat_points(f[f["source"] == s], "fooled", [0, 1], off)
            ax.errorbar([p[0] for p in pts], [p[1] for p in pts], yerr=[p[2] for p in pts],
                        color=COLOR[s], marker="o", markersize=9, lw=2, capsize=4, markeredgecolor="white",
                        markeredgewidth=1.5, label=s.capitalize())
        ax.set_xticks([0, 1], ["Caught the\nwrong advice", "Followed the\nwrong advice"]); ax.set_xlim(-0.5, 1.5)
        ax.set_ylabel("Trust (1 to 5)"); ax.set_ylim(TRUST_MIN - 0.4, TRUST_MAX + 0.4)
        ax.set_title("Trust vs being fooled")
        ax.legend(handles=legend_handles(SOURCES), loc="lower left", ncol=2)
        ax = axes[1]
        oa = c5.dropna(subset=["own_accuracy"]).copy()
        oa["oa"] = oa["own_accuracy"].round(2)
        vals = sorted(oa["oa"].unique())
        for s, off in zip(SOURCES, (-0.06, 0.06)):
            pts = cat_points(oa[oa["source"] == s], "oa", vals, off)
            ax.errorbar([p[0] for p in pts], [p[1] for p in pts], yerr=[p[2] for p in pts],
                        color=COLOR[s], marker="o", markersize=8, lw=2, capsize=3, markeredgecolor="white",
                        markeredgewidth=1.5)
        ax.set_xticks(vals, [f"{100 * v:.0f}%" for v in vals]); ax.set_xlabel("Participant's own share correct on that source's questions")
        ax.set_title("Trust vs own correctness")
        note(fig, "Mean trust (shown when 3+ cells), 95% CI when 8+ cells (rough guide). Both panels share the same scale.")
        fig.tight_layout(rect=(0, 0.04, 1, 1)); save(fig, "5_trust_vs_correctness.png")

    # 6: follow rate by level and source
    fl = per_participant(df, "followed", ["level", "source"])
    fig, ax = plt.subplots(figsize=(7, 4.2))
    xs = np.arange(3); w = 0.36
    for i, s in enumerate(SOURCES):
        vals = [mean_ci(fl[(l, s)]) if (l, s) in fl else (np.nan, 0) for l in LEVELS]
        m = np.array([v[0] for v in vals]) * 100; c = np.array([v[1] for v in vals]) * 100
        pos = xs + (i - 0.5) * (w + 0.02)
        ax.bar(pos, np.nan_to_num(m), w, color=COLOR[s], label=s.capitalize())
        ax.errorbar(pos, np.nan_to_num(m), yerr=c, fmt="none", ecolor=INK_2, elinewidth=1, capsize=3)
        for px, mv in zip(pos, m):
            if pd.notna(mv):
                ax.text(px, 2, f"{mv:.0f}%", ha="center", va="bottom", color="white", fontsize=9, fontweight="bold")
    ax.set_xticks(xs, [l.capitalize() for l in LEVELS]); ax.set_ylim(0, 108)
    ax.set_ylabel("Advice followed (%)"); ax.set_title("How often participants followed the advice")
    ax.legend(loc="lower right", bbox_to_anchor=(1, 1.0), ncol=2)
    note(fig, "Bars: mean across participants. Lines: 95% confidence interval.")
    fig.tight_layout(rect=(0, 0.03, 1, 1)); save(fig, "6_follow_rate_by_level_and_source.png")

    # 7: follow after error vs baseline
    base, post_src = x["base"], x["post_src"]
    cats = ["Baseline", "After robot error", "After screen error"]
    vals = [mean_ci(base)] + [mean_ci(post_src[s]) for s in SOURCES]
    if any(pd.notna(v[0]) for v in vals):
        fig, ax = plt.subplots(figsize=(6.2, 4.2))
        xx = np.arange(3)
        m = [100 * v[0] if pd.notna(v[0]) else 0 for v in vals]
        ax.bar(xx, m, 0.55, color=["#9aa5b1", COLOR["robot"], COLOR["screen"]])
        ax.errorbar(xx, m, yerr=[100 * v[1] for v in vals], fmt="none", ecolor=INK_2, elinewidth=1, capsize=3)
        for px, v in zip(xx, vals):
            if pd.notna(v[0]):
                ax.text(px, 2, f"{100 * v[0]:.0f}%", ha="center", va="bottom", color="white", fontsize=9, fontweight="bold")
        ax.set_xticks(xx, cats); ax.set_ylim(0, 108); ax.set_ylabel("Correct advice followed (%)")
        ax.set_title("Following advice right after a wrong advice")
        note(fig, f"Baseline: correct advice, question {BASELINE_MIN_Q} or later, not right after an error.")
        fig.tight_layout(rect=(0, 0.03, 1, 1)); save(fig, "7_follow_after_error.png")

    # 8: difficulty checks (two panels, different measures)
    fig, axes = plt.subplots(1, 2, figsize=(9, 4))
    acc = x["acc_lvl"]
    m, c = zip(*[mean_ci(acc[l]) for l in LEVELS])
    axes[0].bar(xs, np.nan_to_num(np.array(m)) * 100, 0.55, color=SINGLE)
    axes[0].errorbar(xs, np.nan_to_num(np.array(m)) * 100, yerr=np.array(c) * 100, fmt="none", ecolor=INK_2, capsize=3)
    for px, mv in zip(xs, m):
        if pd.notna(mv):
            axes[0].text(px, 2, f"{100 * mv:.0f}%", ha="center", va="bottom", color="white", fontsize=9, fontweight="bold")
    axes[0].axhline(50, color=INK_2, lw=1, ls="--"); axes[0].text(-0.38, 51.5, "chance (50%)", color=INK_2, fontsize=8, ha="left")
    axes[0].set_ylim(0, 108); axes[0].set_ylabel("Correct bowl chosen (%)"); axes[0].set_title("Accuracy per level")
    d_t = df[df["answered"] & df["time_s"].notna()]
    tm = per_participant(d_t, "time_s", ["level"]).reindex(columns=LEVELS)
    if tm.notna().any().any():
        m, c = zip(*[mean_ci(tm[l]) for l in LEVELS])
        m = np.nan_to_num(np.array(m)); c = np.nan_to_num(np.array(c))
        axes[1].bar(xs, m, 0.55, color=SINGLE); axes[1].errorbar(xs, m, yerr=c, fmt="none", ecolor=INK_2, capsize=3)
        for px, mv, cv in zip(xs, m, c):
            axes[1].text(px, mv + cv + m.max() * 0.03, f"{mv:.1f} s", ha="center", color=INK, fontsize=9)
        axes[1].set_ylim(0, (m + c).max() * 1.25); axes[1].set_ylabel("Decision time (s)")
    axes[1].set_title("Decision time per level")
    for ax in axes:
        ax.set_xticks(xs, [l.capitalize() for l in LEVELS])
    note(fig, "Bars: mean across participants. Lines: 95% confidence interval.")
    fig.tight_layout(rect=(0, 0.03, 1, 1)); save(fig, "8_difficulty_checks.png")
    return files


# ==========================================================================
# 7. Output
# ==========================================================================
def write_excel(T, warns, path):
    from openpyxl.styles import Font, PatternFill, Alignment
    T = dict(T)
    T["Data warnings"] = pd.DataFrame({"warning": warns if warns else ["none"]})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, t in T.items():
            t = t.copy()
            for c in t.columns:
                if t[c].dtype.kind == "f":
                    t[c] = t[c].round(4)
            t.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for cell in ws[1]:
                cell.font = Font(name="Arial", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F3A5F")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Arial")
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
            ws.freeze_panes = "A2"


def run(data_raw, trust_raw, final_raw):
    """Check + compute. Returns (R, warns) or exits with the list of errors."""
    errors, warns, data, trust, final = clean_inputs(data_raw, trust_raw, final_raw)
    if errors:
        print("Please fix these problems in the Excel file first:\n")
        print("\n".join(errors))
        sys.exit(1)
    if data.empty:
        sys.exit("The Data sheet has no rows yet.")
    df = derive(data)
    return compute(df, trust, final), warns


def main():
    ap = argparse.ArgumentParser(description="HRI robot trust study analysis")
    ap.add_argument("excel", help="the filled-in data entry Excel file")
    ap.add_argument("--out", default="results", help="output folder (default: results)")
    args = ap.parse_args()

    data, trust, final = read_workbook(args.excel)
    R, warns = run(data, trust, final)
    if warns:
        print("Warnings (the analysis still runs):")
        print("\n".join("  " + w for w in warns))
        print()
    os.makedirs(args.out, exist_ok=True)
    write_excel(R["T"], warns, os.path.join(args.out, "HRI_results.xlsx"))
    files = charts(R, args.out)
    text = build_summary(R, warns)
    with open(os.path.join(args.out, "summary.txt"), "w", encoding="utf-8") as fh:
        fh.write(text)
    print(text)
    print(f"\nSaved to {args.out}: summary.txt, HRI_results.xlsx and {len(files)} charts")


if __name__ == "__main__":
    main()
