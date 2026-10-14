"""
Builds the data entry workbook.

    python make_template.py            -> HRI_data_entry.xlsx  (empty, with dropdowns and an Example sheet)

build_workbook(path, data, trust, final) also fills the three sheets (used for fake data).
Sheets: Data, Trust, Final (read by analyze.py) and Example (how to fill, ignored by analyze.py).
"""
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")
NOTE_FONT = Font(name="Arial", italic=True, color="5B6470")
TITLE_FONT = Font(name="Arial", bold=True, size=13, color="1F3A5F")

SHEETS = {
    "Data": {"cols": ["participant", "level", "question", "source", "advice", "correct", "chosen", "time_s"],
             "widths": [13, 10, 10, 10, 10, 10, 10, 10], "max_row": 1500},
    "Trust": {"cols": ["participant", "level", "trust_robot", "trust_screen", "rated_first"],
              "widths": [13, 10, 13, 13, 13], "max_row": 400},
    "Final": {"cols": ["participant", "more_trusted", "reason"],
              "widths": [13, 14, 60], "max_row": 200},
}


def _list(ws, col_letter, items, last_row, title, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(str(i) for i in items) + '"', allow_blank=True,
                        showErrorMessage=True, errorStyle="stop", errorTitle=title,
                        error="Pick a value from the dropdown list.")
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = title, prompt, True
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")


def _validations(name, ws):
    last = SHEETS[name]["max_row"]
    if name == "Data":
        _list(ws, "B", ["easy", "medium", "hard"], last, "level")
        _list(ws, "C", [1, 2, 3, 4, 5, 6], last, "question")
        _list(ws, "D", ["robot", "screen"], last, "source")
        for col, t in [("E", "advice"), ("F", "correct"), ("G", "chosen")]:
            _list(ws, col, ["red", "blue"], last, t)
        dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorStyle="stop", errorTitle="time_s",
                            error="Type seconds as a number above 0, for example 4.2 (no units).")
        ws.add_data_validation(dv)
        dv.add(f"H2:H{last}")
    elif name == "Trust":
        _list(ws, "B", ["easy", "medium", "hard"], last, "level")
        for col, t in [("C", "trust_robot"), ("D", "trust_screen")]:
            _list(ws, col, [1, 2, 3, 4, 5], last, t, "1 = lowest trust, 5 = highest trust")
        _list(ws, "E", ["robot", "screen"], last, "rated_first", "Which source was the participant asked about first?")
    else:
        _list(ws, "B", ["robot", "screen", "equal"], last, "more_trusted")


def _sheet(wb, name, df):
    ws = wb.create_sheet(name)
    spec = SHEETS[name]
    for j, (c, w) in enumerate(zip(spec["cols"], spec["widths"]), start=1):
        cell = ws.cell(row=1, column=j, value=c)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"
    if df is not None:
        for i, row in enumerate(df[spec["cols"]].itertuples(index=False), start=2):
            for j, v in enumerate(row, start=1):
                if pd.isna(v):
                    continue
                cell = ws.cell(row=i, column=j, value=v.item() if hasattr(v, "item") else v)
                cell.font = BODY_FONT
    _validations(name, ws)
    return ws


def _example(wb):
    """Example sheet with one made-up participant, plus the rules for each column."""
    from make_fake_data import generate
    d, t, f = generate(1, seed=7)
    ws = wb.create_sheet("Example")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "How to fill in this workbook (example with one made-up participant, P01)"
    ws["A1"].font = TITLE_FONT
    notes = [
        "Type into the sheets Data, Trust and Final. This Example sheet is ignored by the analysis. Use lowercase or uppercase, both work.",
        "Use the dropdowns. A value that is not in the list is refused. Type time_s as a number (4.2 or 4,2), no units.",
        "Data: one row per question, 18 rows per participant. advice = bowl the robot pointed to or the screen showed. "
        "correct = right bowl from the answer key. chosen = bowl the participant used.",
        "Trust: one row per participant and level (3 rows per participant). Ratings from 1 (lowest) to 5 (highest). "
        "rated_first = which source was asked about first in that level.",
        "Final: one row per participant. more_trusted = robot, screen or equal. reason = a few words.",
        "Use the same ID everywhere (P01 is P01 in all three sheets). If the analysis stops, it names the sheet and the Excel row to fix.",
        "Not decided yet: the positions of the wrong advices. The analysis only warns about the design, so any positions can be entered.",
    ]
    for i, n in enumerate(notes, start=3):
        ws.cell(row=i, column=1, value=n).font = NOTE_FONT
    r = len(notes) + 4
    for title, df, cols in [("Data sheet (first 8 of 18 rows shown)", d.head(8), SHEETS["Data"]["cols"]),
                            ("Trust sheet", t, SHEETS["Trust"]["cols"]),
                            ("Final sheet", f, SHEETS["Final"]["cols"])]:
        ws.cell(row=r, column=1, value=title).font = Font(name="Arial", bold=True, size=11)
        r += 1
        for j, c in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=j, value=c)
            cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        for row in df[cols].itertuples(index=False):
            r += 1
            for j, v in enumerate(row, start=1):
                ws.cell(row=r, column=j, value=v.item() if hasattr(v, "item") else v).font = BODY_FONT
        r += 3
    for col, w in zip("ABCDEFGH", [13, 12, 14, 14, 14, 10, 10, 10]):
        ws.column_dimensions[col].width = w


def build_workbook(path, data=None, trust=None, final=None):
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Data", data)
    _sheet(wb, "Trust", trust)
    _sheet(wb, "Final", final)
    _example(wb)
    wb.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "HRI_data_entry.xlsx"
    build_workbook(out)
    print(f"Wrote {out}")
